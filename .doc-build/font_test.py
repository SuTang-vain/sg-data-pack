from docx import Document
from pathlib import Path
import os
from docx.shared import Pt
from docx.oxml.ns import qn
fonts=['华文黑体','黑体-简','冬青黑体简体中文','STHeiti','Heiti SC','Hiragino Sans GB','Arial Unicode MS','Songti SC','PingFang SC','SimSun','Microsoft YaHei','Arial Unicode']
d=Document()
for f in fonts:
 p=d.add_paragraph(); r=p.add_run(f+'：组件库数据治理与可信变更链 你好世界'); r.font.name=f; r.font.size=Pt(18); r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),f); r._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),f); r._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),f)
out = Path(os.environ.get('SG_DOC_OUTPUT_DIR', str(Path(__file__).resolve().parent / 'output'))).expanduser().resolve()
out.mkdir(parents=True, exist_ok=True)
d.save(out / 'font-test.docx')
